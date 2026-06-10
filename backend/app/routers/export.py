from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from app.config import supabase
from app.dependencies import require_admin
from app.models.schemas import ExportRequest
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

router = APIRouter(prefix="/export", tags=["Export"])


def get_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def fmt_date(val):
    if not val or val == "-":
        return "-"
    try:
        return datetime.fromisoformat(str(val).replace("Z", "+00:00")).strftime("%d/%m/%Y %H:%M")
    except:
        return "-"


def fmt_date_only(val):
    if not val or val == "-":
        return "-"
    try:
        return datetime.fromisoformat(str(val).replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except:
        return "-"


ROLE_ORDER = {
    "main_leader": 1,
    "assistant_leader": 2,
    "floater": 3,
    "operator": 4,
}

ROLE_COLORS = {
    "main_leader": "6366F1",
    "assistant_leader": "06B6D4",
    "floater": "10B981",
    "operator": "F59E0B",
}


def build_workbook(ws, title, subtitle, headers, col_widths, rows_data, row_builder):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    merge_end = chr(64 + len(headers))
    ws.merge_cells(f"A1:{merge_end}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1E293B")
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{merge_end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color="64748B")
    ws["A2"].alignment = center
    ws.row_dimensions[2].height = 20

    for col, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = get_border()
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[4].height = 22

    for i, row in enumerate(rows_data, 1):
        row_data, fill_color, is_italic, name_col = row_builder(i, row)
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=i + 4, column=col, value=value)
            cell.alignment = left if col == name_col else center
            cell.border = get_border()
            cell.fill = row_fill
            if is_italic:
                cell.font = Font(color="1E293B", italic=True)
        ws.row_dimensions[i + 4].height = 18


@router.post("/excel")
def export_excel(body: ExportRequest, current_user: dict = Depends(require_admin)):
    try:
        report_type = getattr(body, 'report_type', 'active')
        wb = openpyxl.Workbook()
        ws = wb.active
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        month_label = body.month_year or "All Months"

        if report_type == 'removed':
            ws.title = "Removed Users Report"

            history_query = supabase.table("user_history").select("*")
            if body.shift:
                history_query = history_query.eq("shift", body.shift)
            if body.month_year:
                history_query = history_query.eq("month_year", body.month_year)
            if body.project_id:
                proj = supabase.table("projects").select("project_number").eq(
                    "id", body.project_id
                ).single().execute()
                if proj.data:
                    history_query = history_query.eq("project_number", proj.data["project_number"])

            history_result = history_query.execute()
            raw = history_result.data or []

            rows = sorted(raw, key=lambda x: (
                x.get("project_number", ""),
                x.get("line_type", ""),
                x.get("shift", ""),
                ROLE_ORDER.get(x.get("role"), 99),
                x.get("removed_at", ""),
            ))

            headers    = ["No.", "Project", "Line", "Station", "Shift", "Employee ID", "Badge/PIN", "Full Name", "Role", "Registered", "Removed At"]
            col_widths = [6,      12,         14,      12,        12,      18,            24,           28,          20,     16,            22]

            def removed_row_builder(i, h):
                role = h.get("role", "-")
                row_data = [
                    i,
                    h.get("project_number", "-"),
                    h.get("line_type", "-"),
                    h.get("station_code", "-"),
                    h.get("shift", "-"),
                    h.get("employee_id", "-"),
                    h.get("badge_number") or "-",
                    h.get("full_name", "-"),
                    role.replace("_", " ").title(),
                    fmt_date_only(h.get("registered_at")),
                    fmt_date(h.get("removed_at")),
                ]
                return row_data, "FEF2F2", True, 8

            build_workbook(
                ws,
                f"MES Removed Users Report — {month_label}",
                f"Generated: {now_str}  |  V.S. Industry Berhad  |  Total: {len(rows)} records",
                headers, col_widths, rows, removed_row_builder
            )

            summary_row = len(rows) + 6
            ws.merge_cells(f"A{summary_row}:C{summary_row}")
            ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True, size=12)
            summary_data = [
                ("Total Removed",      len(rows)),
                ("Main Leaders",       sum(1 for r in rows if r.get("role") == "main_leader")),
                ("Assistant Leaders",  sum(1 for r in rows if r.get("role") == "assistant_leader")),
                ("Floaters",           sum(1 for r in rows if r.get("role") == "floater")),
                ("Operators",          sum(1 for r in rows if r.get("role") == "operator")),
            ]
            for j, (label, count) in enumerate(summary_data, 1):
                ws.cell(row=summary_row + j, column=1, value=label).font = Font(bold=True)
                ws.cell(row=summary_row + j, column=2, value=count).font = Font(bold=True, color="EF4444")

        else:
            ws.title = "Active Users Report"
            rows = []

            users_query = supabase.table("users").select(
                "id, employee_id, full_name, role, shift, badge_number, created_at, "
                "projects(project_number, name), lines(line_type)"
            ).neq("status", "inactive").neq("role", "mes_admin").neq("role", "operator")

            if body.project_id:
                users_query = users_query.eq("project_id", body.project_id)
            if body.line_id:
                users_query = users_query.eq("line_id", body.line_id)
            if body.shift:
                users_query = users_query.eq("shift", body.shift)

            leaders_result = users_query.execute()
            for u in (leaders_result.data or []):
                project = u.get("projects") or {}
                line = u.get("lines") or {}
                rows.append({
                    "project_number": project.get("project_number", "-"),
                    "line_type": line.get("line_type", "-"),
                    "station_code": "-",
                    "shift": u.get("shift", "-"),
                    "employee_id": u.get("employee_id", "-"),
                    "badge_number": u.get("badge_number") or "-",
                    "full_name": u.get("full_name", "-"),
                    "role": u.get("role", "-"),
                    "role_order": ROLE_ORDER.get(u.get("role"), 99),
                    "registered_at": u.get("created_at"),
                })

            assign_query = supabase.table("station_assignments").select(
                "shift, month_year, status, "
                "snapshot_name, snapshot_employee_id, snapshot_role, snapshot_badge_number, "
                "stations!station_assignments_station_id_fkey(station_code, line_id, lines(line_type, project_id, projects(project_number, name))), "
                "users!station_assignments_user_id_fkey(id, full_name, employee_id, role, badge_number, created_at)"
            ).eq("status", "active")

            if body.month_year:
                assign_query = assign_query.eq("month_year", body.month_year)
            if body.shift:
                assign_query = assign_query.eq("shift", body.shift)

            assign_result = assign_query.execute()
            assignments = assign_result.data or []

            if body.project_id:
                assignments = [
                    a for a in assignments
                    if a.get("stations", {}).get("lines", {}).get("project_id") == body.project_id
                ]
            if body.line_id:
                assignments = [
                    a for a in assignments
                    if a.get("stations", {}).get("line_id") == body.line_id
                ]

            seen_active = set()
            for a in assignments:
                user = a.get("users") or {}
                role = user.get("role", "operator")
                if role != "operator":
                    continue
                uid = user.get("id")
                if uid in seen_active:
                    continue
                seen_active.add(uid)
                station = a.get("stations") or {}
                line = station.get("lines") or {}
                project = line.get("projects") or {}
                rows.append({
                    "project_number": project.get("project_number", "-"),
                    "line_type": line.get("line_type", "-"),
                    "station_code": station.get("station_code", "-"),
                    "shift": a.get("shift", "-"),
                    "employee_id": user.get("employee_id", "-"),
                    "badge_number": user.get("badge_number") or "-",
                    "full_name": user.get("full_name", "-"),
                    "role": role,
                    "role_order": ROLE_ORDER.get(role, 99),
                    "registered_at": user.get("created_at"),
                })

            rows.sort(key=lambda x: (
                x.get("project_number", ""),
                x.get("line_type", ""),
                x.get("shift", ""),
                x.get("role_order", 99),
                x.get("full_name", ""),
            ))

            headers    = ["No.", "Project", "Line", "Station", "Shift", "Employee ID", "Badge/PIN", "Full Name", "Role", "Registered", "Status"]
            col_widths = [6,      12,         14,      12,        12,      18,            24,           28,          20,     16,            14]

            def active_row_builder(i, row):
                role = row.get("role", "operator")
                color = ROLE_COLORS.get(role, "94A3B8")
                row_data = [
                    i,
                    row.get("project_number", "-"),
                    row.get("line_type", "-"),
                    row.get("station_code", "-"),
                    row.get("shift", "-"),
                    row.get("employee_id", "-"),
                    row.get("badge_number", "-"),
                    row.get("full_name", "-"),
                    role.replace("_", " ").title(),
                    fmt_date_only(row.get("registered_at")),
                    "Active",
                ]
                return row_data, color, False, 8

            build_workbook(
                ws,
                f"MES Active Users Report — {month_label}",
                f"Generated: {now_str}  |  V.S. Industry Berhad  |  Total: {len(rows)} records",
                headers, col_widths, rows, active_row_builder
            )

            summary_row = len(rows) + 6
            ws.merge_cells(f"A{summary_row}:C{summary_row}")
            ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True, size=12)
            summary_data = [
                ("Total Active",       len(rows)),
                ("Main Leaders",       sum(1 for r in rows if r["role"] == "main_leader")),
                ("Assistant Leaders",  sum(1 for r in rows if r["role"] == "assistant_leader")),
                ("Floaters",           sum(1 for r in rows if r["role"] == "floater")),
                ("Operators",          sum(1 for r in rows if r["role"] == "operator")),
            ]
            for j, (label, count) in enumerate(summary_data, 1):
                ws.cell(row=summary_row + j, column=1, value=label).font = Font(bold=True)
                ws.cell(row=summary_row + j, column=2, value=count).font = Font(bold=True, color="6366F1")

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        rtype = "Removed" if report_type == "removed" else "Active"
        filename = f"MES_{rtype}_Report_{month_label}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except Exception as e:
        print(f"EXPORT ERROR: {str(e)}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")


@router.get("/history")
def get_export_history(current_user: dict = Depends(require_admin)):
    result = supabase.table("station_assignments").select(
        "month_year"
    ).in_("status", ["active", "expired"]).execute()

    months = list(set([r["month_year"] for r in result.data or [] if r.get("month_year")]))
    months.sort(reverse=True)
    return {"months": months}